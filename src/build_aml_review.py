from __future__ import annotations

from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# =========================================================
# 3. AML Review Workflow
# =========================================================
# Build a player-level AML review file from the Q1 player-day result,
# then validate the returned workbook and export the CSV replica.

PROJECT_ROOT = Path(__file__).resolve().parents[1]
Q1_RESULT_XLSX = PROJECT_ROOT / "output" / "player_day.xlsx"
Q3_OUTPUT_XLSX = PROJECT_ROOT / "output" / "aml_review.xlsx"
Q3_OUTPUT_CSV = PROJECT_ROOT / "output" / "aml_review.csv"
Q3_RETURNED_XLSX = PROJECT_ROOT / "output" / "aml_review_returned.xlsx"

EXPECTED_COLUMNS = [
    "Activity Date",
    "Player ID",
    "Gender",
    "Age",
    "Country",
    "Wager Amount",
    "Winning Amount",
    "Deposit Amount",
    "Withdrawal Amount",
    "Review Status",
    "Decision",
    "Comment",
]

ALLOWED_REVIEW_STATUS = [
    "To be reviewed",
    "Review in Progress",
    "Review Completed",
]
ALLOWED_DECISIONS = [
    "No Action required",
    "Operator Review needed",
    "Account Freeze requested",
]

Q1_REQUIRED_COLUMNS = [
    "UserID",
    "ActivityDate",
    "Gender",
    "Age",
    "Country",
    "CashGameWagers",
    "CashGameWinnings",
    "TournamentWagers",
    "TournamentWinnings",
    "DepositAmount",
    "WithdrawalAmount",
]

# =========================================================
# 3.1 Load and validate the Q1 result
# =========================================================

def load_q1_result(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Missing Q1 output: {path}")

    df = pd.read_excel(path, usecols=Q1_REQUIRED_COLUMNS)
    missing_columns = sorted(set(Q1_REQUIRED_COLUMNS) - set(df.columns))
    if missing_columns:
        raise ValueError(f"Q1 output is missing required columns: {missing_columns}")

    if df.empty:
        raise ValueError("Q1 output contains no player-day records.")

    df["ActivityDate"] = pd.to_datetime(df["ActivityDate"], errors="coerce")
    if df["ActivityDate"].isna().any():
        raise ValueError("Q1 output contains invalid ActivityDate values.")

    df["UserID"] = pd.to_numeric(df["UserID"], errors="coerce")
    if df["UserID"].isna().any():
        raise ValueError("Q1 output contains invalid UserID values.")
    df["UserID"] = df["UserID"].astype("int64")

    numeric_columns = [
        "Age",
        "CashGameWagers",
        "CashGameWinnings",
        "TournamentWagers",
        "TournamentWinnings",
        "DepositAmount",
        "WithdrawalAmount",
    ]
    for column in numeric_columns:
        df[column] = pd.to_numeric(df[column], errors="coerce")

    invalid_age = df["Age"].isna()
    if invalid_age.any():
        raise ValueError("Q1 output contains invalid Age values.")

    amount_columns = [column for column in numeric_columns if column != "Age"]
    df[amount_columns] = df[amount_columns].fillna(0)
    df["Gender"] = df["Gender"].astype("string").str.strip()
    df["Country"] = df["Country"].astype("string").str.strip()

    return df

# =========================================================
# 3.2 Select Top 50 players
# =========================================================

def build_top_50_players(q1_df: pd.DataFrame) -> pd.DataFrame:
    # Assumption: for a player-level AML review, the relevant activity date is the player's
    # latest successful withdrawal date. This is the most stable business interpretation of
    # the required "Activity Date" column for the top 50 accounts.
    working = q1_df.copy()
    working["WagerAmount"] = working["CashGameWagers"] + working["TournamentWagers"]
    working["WinningAmount"] = working["CashGameWinnings"] + working["TournamentWinnings"]

    player_totals = (
        working.groupby("UserID", as_index=False)
        .agg(
            WagerAmount=("WagerAmount", "sum"),
            WinningAmount=("WinningAmount", "sum"),
            DepositAmount=("DepositAmount", "sum"),
            WithdrawalAmount=("WithdrawalAmount", "sum"),
        )
    )

    latest_withdrawal_dates = (
        working.loc[working["WithdrawalAmount"] > 0, ["UserID", "ActivityDate"]]
        .groupby("UserID", as_index=False)
        .agg(ActivityDate=("ActivityDate", "max"))
    )

    latest_player_details = (
        working.merge(latest_withdrawal_dates, on=["UserID", "ActivityDate"], how="inner")
        [["UserID", "ActivityDate", "Gender", "Age", "Country"]]
        .drop_duplicates(subset=["UserID"])
    )

    player_summary = latest_player_details.merge(player_totals, on="UserID", how="inner")
    player_summary = (
        player_summary.sort_values(
            ["WithdrawalAmount", "UserID"],
            ascending=[False, True],
        )
        .head(50)
        .copy()
    )

    if len(player_summary) != 50:
        raise ValueError(
            f"Expected 50 players with successful withdrawals, found {len(player_summary)}."
        )

    player_summary["ActivityDate"] = player_summary["ActivityDate"].dt.date
    player_summary["Age"] = player_summary["Age"].astype(int)

    return player_summary.rename(columns={"UserID": "PlayerID"})

# =========================================================
# 3.3 Create AML review workbook
# =========================================================

def build_aml_table(player_df: pd.DataFrame) -> pd.DataFrame:
    result = pd.DataFrame({
        "Activity Date": player_df["ActivityDate"],
        "Player ID": player_df["PlayerID"].astype(str),
        "Gender": player_df["Gender"],
        "Age": player_df["Age"],
        "Country": player_df["Country"].astype("string").str.strip(),
        "Wager Amount": player_df["WagerAmount"],
        "Winning Amount": player_df["WinningAmount"],
        "Deposit Amount": player_df["DepositAmount"],
        "Withdrawal Amount": player_df["WithdrawalAmount"],
        "Review Status": "To be reviewed",
        "Decision": "",
        "Comment": "",
    })
    return result[EXPECTED_COLUMNS]


def style_aml_sheet(ws, n_rows: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 24

    for row in ws.iter_rows(min_row=2, max_row=1 + n_rows, min_col=1, max_col=12):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    # Lock first 9 columns; leave last 3 editable.
    for row in ws.iter_rows(min_row=2, max_row=1 + n_rows, min_col=1, max_col=9):
        for cell in row:
            prot = copy(cell.protection)
            prot.locked = True
            cell.protection = prot

    for row in ws.iter_rows(min_row=2, max_row=1 + n_rows, min_col=10, max_col=12):
        for cell in row:
            prot = copy(cell.protection)
            prot.locked = False
            cell.protection = prot

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{1 + n_rows}"

    column_widths = {
        "Activity Date": 14,
        "Player ID": 12,
        "Gender": 10,
        "Age": 8,
        "Country": 16,
        "Wager Amount": 14,
        "Winning Amount": 15,
        "Deposit Amount": 15,
        "Withdrawal Amount": 17,
        "Review Status": 18,
        "Decision": 24,
        "Comment": 32,
    }
    for idx, col in enumerate(EXPECTED_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = column_widths[col]

    for row in range(2, n_rows + 2):
        ws.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=2).number_format = "@"
        ws.cell(row=row, column=4).number_format = "0"
        for column in range(6, 10):
            ws.cell(row=row, column=column).number_format = "#,##0.00"

    ws.protection.sheet = True
    ws.protection.enable()
    ws.sheet_view.showGridLines = False

    status_dv = DataValidation(
        type="list",
        formula1='"To be reviewed,Review in Progress,Review Completed"',
        allow_blank=False,
        showErrorMessage=True,
        errorStyle="stop",
        showInputMessage=True,
    )
    status_dv.errorTitle = "Invalid review status"
    status_dv.error = "Please choose a valid review status."
    status_dv.promptTitle = "Review status"
    status_dv.prompt = "Choose one of the allowed review status values."
    status_dv.add(f"J2:J{1 + n_rows}")
    ws.add_data_validation(status_dv)

    decision_dv = DataValidation(
        type="list",
        formula1='"No Action required,Operator Review needed,Account Freeze requested"',
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        showInputMessage=True,
    )
    decision_dv.errorTitle = "Invalid decision"
    decision_dv.error = "Please choose a valid decision."
    decision_dv.promptTitle = "AML decision"
    decision_dv.prompt = "Choose one of the allowed decisions."
    decision_dv.add(f"K2:K{1 + n_rows}")
    ws.add_data_validation(decision_dv)

# =========================================================
# 3.4 Save workbook with locked / editable columns
# =========================================================

def write_aml_workbook(df: pd.DataFrame, out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "AML Review"
    ws.append(EXPECTED_COLUMNS)
    for row in df.itertuples(index=False, name=None):
        ws.append(row)
    style_aml_sheet(ws, len(df))
    wb.save(out_path)

# =========================================================
# 3.5 Create an optional demonstration return file
# =========================================================

def create_demo_aml_return(source_path: Path, target_path: Path, baseline_df: pd.DataFrame) -> None:
    """Create synthetic AML inputs only when no returned workbook is available."""
    wb = load_workbook(source_path)
    ws = wb["AML Review"]

    status_values = [
        "Review in Progress",
        "Review Completed",
        "To be reviewed",
    ]
    decision_values = [
        "Operator Review needed",
        "No Action required",
        "Account Freeze requested",
    ]

    for idx, row in enumerate(range(2, len(baseline_df) + 2), start=0):
        ws.cell(row=row, column=10).value = status_values[idx % len(status_values)]
        ws.cell(row=row, column=11).value = decision_values[idx % len(decision_values)]
        ws.cell(row=row, column=12).value = "Synthetic input for workflow demonstration only."

    wb.save(target_path)

# =========================================================
# 3.6 Validate returned workbook and export CSV replica
# =========================================================

def read_returned_workbook(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Missing returned AML workbook: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "AML Review" not in workbook.sheetnames:
        raise ValueError("Returned workbook is missing the 'AML Review' worksheet.")

    return pd.read_excel(
        path,
        sheet_name="AML Review",
        dtype={"Player ID": str},
    )


def normalize_first_nine_columns(df: pd.DataFrame) -> pd.DataFrame:
    normalized = df[EXPECTED_COLUMNS[:9]].copy().reset_index(drop=True)
    normalized["Activity Date"] = pd.to_datetime(
        normalized["Activity Date"],
        errors="coerce",
    ).dt.date
    if normalized["Activity Date"].isna().any():
        raise ValueError("Activity Date contains an invalid date.")

    normalized["Player ID"] = normalized["Player ID"].astype("string").str.strip()
    normalized["Gender"] = normalized["Gender"].astype("string").str.strip()
    normalized["Country"] = normalized["Country"].astype("string").str.strip()

    numeric_columns = [
        "Age",
        "Wager Amount",
        "Winning Amount",
        "Deposit Amount",
        "Withdrawal Amount",
    ]
    for column in numeric_columns:
        normalized[column] = pd.to_numeric(normalized[column], errors="coerce")
        if normalized[column].isna().any():
            raise ValueError(f"{column} contains an invalid numeric value.")

    normalized["Age"] = normalized["Age"].astype(int)
    return normalized


def validate_returned_workbook(returned_df: pd.DataFrame, baseline_df: pd.DataFrame) -> None:
    if list(returned_df.columns) != EXPECTED_COLUMNS:
        raise ValueError("Columns are not in the expected order.")
    if len(returned_df) != len(baseline_df) or len(returned_df) != 50:
        raise ValueError("Workbook does not contain exactly 50 player records.")
    if not returned_df["Player ID"].is_unique:
        raise ValueError("Player ID is not unique.")

    returned_first_nine = normalize_first_nine_columns(returned_df)
    baseline_first_nine = normalize_first_nine_columns(baseline_df)

    try:
        pd.testing.assert_frame_equal(
            returned_first_nine,
            baseline_first_nine,
            check_dtype=False,
            rtol=0,
            atol=1e-8,
        )
    except AssertionError as exc:
        raise ValueError(
            "One or more protected source fields were changed or misplaced."
        ) from exc

    allowed_review = set(ALLOWED_REVIEW_STATUS)
    allowed_decisions = set(ALLOWED_DECISIONS)

    review_status = returned_df["Review Status"].astype("string").str.strip()
    decision = returned_df["Decision"].astype("string").str.strip()

    if review_status.isna().any():
        raise ValueError("Review Status contains a missing value.")
    invalid_review_status = sorted(set(review_status) - allowed_review)
    if invalid_review_status:
        raise ValueError(f"Invalid Review Status values: {invalid_review_status}")

    if decision.isna().any():
        raise ValueError("Decision contains a missing value.")
    invalid_decisions = sorted(set(decision) - allowed_decisions)
    if invalid_decisions:
        raise ValueError(f"Invalid Decision values: {invalid_decisions}")

    returned_df["Review Status"] = review_status
    returned_df["Decision"] = decision
    returned_df["Comment"] = returned_df["Comment"].fillna("").astype(str)


def export_csv_replica(returned_df: pd.DataFrame, output_path: Path) -> None:
    csv_df = returned_df.copy()
    csv_df["Activity Date"] = pd.to_datetime(
        csv_df["Activity Date"],
        errors="raise",
    ).dt.strftime("%Y-%m-%d")
    csv_df.to_csv(output_path, index=False)

# =========================================================
# 3.7 Main workflow
# =========================================================

def main() -> None:
    q1_df = load_q1_result(Q1_RESULT_XLSX)
    player_df = build_top_50_players(q1_df)
    aml_df = build_aml_table(player_df)
    write_aml_workbook(aml_df, Q3_OUTPUT_XLSX)

    created_demo_return = False
    if not Q3_RETURNED_XLSX.exists():
        create_demo_aml_return(Q3_OUTPUT_XLSX, Q3_RETURNED_XLSX, aml_df)
        created_demo_return = True

    returned_df = read_returned_workbook(Q3_RETURNED_XLSX)
    validate_returned_workbook(returned_df, aml_df)
    export_csv_replica(returned_df, Q3_OUTPUT_CSV)

    print(f"Saved AML review workbook: {Q3_OUTPUT_XLSX}")
    print(f"Validated returned workbook: {Q3_RETURNED_XLSX}")
    print(f"Saved CSV replica: {Q3_OUTPUT_CSV}")
    if created_demo_return:
        print("The returned workbook contains synthetic demonstration inputs only.")


if __name__ == "__main__":
    main()
